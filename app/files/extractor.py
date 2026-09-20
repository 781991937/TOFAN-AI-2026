"""Safe extraction of teaching text from supported student files."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class FileExtractionError(ValueError):
    pass


def extract_teaching_text(filename: str, data: bytes) -> tuple[str, int | None]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise FileExtractionError("Only PDF, DOCX, and TXT files are supported.")

    try:
        if suffix == ".txt":
            return data.decode("utf-8-sig"), 1

        if suffix == ".pdf":
            reader = PdfReader(BytesIO(data))
            pages = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")
            return "\n\n".join(pages).strip(), len(reader.pages)

        document = Document(BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        tables = []
        for table in document.tables:
            for row in table.rows:
                tables.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs + tables).strip(), None
    except Exception as exc:
        raise FileExtractionError("The file could not be read or extracted.") from exc
