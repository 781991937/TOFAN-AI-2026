from pathlib import Path
import json
import re
import unicodedata

from docx import Document
from pypdf import PdfReader

try:
    import pdftotext
except ImportError:  # pragma: no cover
    pdftotext = None

try:
    import pymupdf
except ImportError:  # pragma: no cover
    pymupdf = None

PAGE_MARKER_RE = re.compile(r"^\s*\[\[PAGE:(\d+)\]\]\s*$")
PAGE_MARKER_ANY_RE = re.compile(r"\[\[PAGE:(\d+)\]\]")
LESSON_MARKER_RE = re.compile(r"^\s*\[\[LESSON:(.*?)\]\]\s*$", re.I)

_ARABIC_ORDINALS = {
    "الأول": 1, "الاول": 1, "أول": 1, "اول": 1,
    "الثاني": 2, "الثانى": 2, "ثاني": 2, "ثانى": 2,
    "الثالث": 3, "ثالث": 3,
    "الرابع": 4, "رابع": 4,
    "الخامس": 5, "خامس": 5,
    "السادس": 6, "سادس": 6,
    "السابع": 7, "سابع": 7,
    "الثامن": 8, "ثامن": 8,
    "التاسع": 9, "تاسع": 9,
    "العاشر": 10, "عاشر": 10,
}
_ENGLISH_ORDINALS = {
    "first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5,
    "sixth": 6, "seventh": 7, "eighth": 8, "ninth": 9, "tenth": 10,
}


def _arabic_normalize_for_matching(value: str) -> str:
    value = unicodedata.normalize("NFC", value or "")
    # Normalize common OCR/PDF variants only for heading detection.
    return (
        value.replace("إ", "ا").replace("أ", "ا").replace("آ", "ا")
        .replace("ٱ", "ا").replace("ـ", "")
    )


class FileExtractor:
    SUPPORTED = {
        ".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".html", ".css",
        ".xml", ".yaml", ".yml", ".rst", ".tex", ".log", ".ini", ".cfg", ".rtf", ".xlsx", ".pptx"
    }

    def extract(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix not in self.SUPPORTED:
            raise ValueError(f"Unsupported file type: {suffix}")
        if suffix == ".pdf":
            return self._pdf(path)
        if suffix == ".docx":
            return self._docx(path)
        if suffix == ".xlsx":
            return self._xlsx(path)
        if suffix == ".pptx":
            return self._pptx(path)
        if suffix == ".json":
            try:
                return json.dumps(json.loads(path.read_text(encoding="utf-8", errors="replace")), ensure_ascii=False, indent=2)
            except Exception:
                pass
        text = path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".rtf":
            text = re.sub(r"\\[a-z]+\d* ?|[{}]", "", text)
        return text

    @staticmethod
    def _pdf(path: Path) -> str:
        toc_by_page: dict[int, list[str]] = {}
        if pymupdf is not None:
            try:
                doc = pymupdf.open(str(path))
                for item in doc.get_toc():
                    if len(item) >= 3 and int(item[0]) == 1 and int(item[2]) >= 1:
                        title = str(item[1] or "").strip()
                        if title:
                            toc_by_page.setdefault(int(item[2]), []).append(title)
                doc.close()
            except Exception:
                toc_by_page = {}

        if pdftotext is not None:
            try:
                with path.open("rb") as handle:
                    pdf = pdftotext.PDF(handle)
                parts = []
                for number, page in enumerate(pdf, 1):
                    for title in toc_by_page.get(number, []):
                        parts.append(f"[[LESSON:{title}]]")
                    parts.extend([f"[[PAGE:{number}]]", page or ""])
                return "\n\n".join(parts)
            except Exception:
                pass

        reader = PdfReader(str(path))
        parts = []
        for number, page in enumerate(reader.pages, 1):
            for title in toc_by_page.get(number, []):
                parts.append(f"[[LESSON:{title}]]")
            parts.extend([f"[[PAGE:{number}]]", page.extract_text() or ""])
        return "\n\n".join(parts)

    @staticmethod
    def _docx(path: Path) -> str:
        document = Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)

    @staticmethod
    def _xlsx(path: Path) -> str:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[ورقة: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                if values:
                    parts.append(" | ".join(values))
        return "\n".join(parts)

    @staticmethod
    def _pptx(path: Path) -> str:
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"[[PAGE:{i}]]")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
        return "\n".join(parts)


def clean_text(text: str) -> str:
    """Clean extracted text without reversing Arabic or other RTL languages.

    PDF extraction may already return logical RTL text. Reversing Arabic words
    here corrupts the source, so cleanup is deliberately Unicode-safe and
    language-neutral.
    """
    text = (text or "").replace("\x00", "").replace("\u00ad", "").replace("\ufeff", "")
    text = unicodedata.normalize("NFC", text)
    cleaned = []
    for raw in text.replace("\f", "\n").splitlines():
        line = " ".join(raw.split()).strip()
        if not line:
            continue
        page = PAGE_MARKER_ANY_RE.fullmatch(line)
        if page:
            cleaned.append(f"[[PAGE:{int(page.group(1))}]]")
            continue
        lesson = LESSON_MARKER_RE.fullmatch(line)
        if lesson:
            title = lesson.group(1).strip()
            if title:
                cleaned.append(f"[[LESSON:{title}]]")
            continue
        if re.fullmatch(r"(?:Page|صفحة)\s*\d+", line, re.I):
            continue
        if re.fullmatch(r"[-_=·•\s]{3,}", line):
            continue
        cleaned.append(line)
    return "\n".join(cleaned).strip()


def page_parts(text: str, chars_per_page: int = 1800) -> list[tuple[int, str]]:
    text = clean_text(text)
    if not text:
        return []
    matches = list(PAGE_MARKER_RE.finditer(text))
    if matches:
        pages = []
        for i, match in enumerate(matches):
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            body = text[match.end():end].strip()
            body = re.sub(r"\[\[LESSON:.*?\]\]\s*", "", body, flags=re.I)
            if body:
                pages.append((int(match.group(1)), body))
        return pages

    chunks, current, size, page_no = [], [], 0, 1
    for line in text.splitlines():
        if current and size + len(line) + 1 > chars_per_page:
            chunks.append((page_no, "\n".join(current).strip()))
            page_no += 1
            current, size = [], 0
        current.append(line)
        size += len(line) + 1
    if current:
        chunks.append((page_no, "\n".join(current).strip()))
    return chunks


def _lesson_heading(line: str):
    """Recognize Arabic/English lesson headings, including common PDF OCR variants."""
    value = line.strip()
    if not value or len(value) > 180:
        return None

    normalized = _arabic_normalize_for_matching(value)
    ordinal_pattern = "|".join(re.escape(x) for x in _ARABIC_ORDINALS)
    arabic = re.match(
        rf"^\s*(?:الدرس|درس|المحاضرة|محاضرة|الوحدة|وحدة|الفصل|فصل)\s*(?:رقم\s*)?({ordinal_pattern}|[0-9٠-٩]+)\b\s*[:：\-–—.]?\s*(.*)$",
        normalized,
        re.I,
    )
    if arabic:
        original_tail = value[value.lower().find(arabic.group(2).lower()):] if arabic.group(2) else ""
        return (original_tail or arabic.group(2) or value).strip()

    english_ordinal = "|".join(_ENGLISH_ORDINALS)
    english = re.match(
        rf"^\s*(?:lesson|lecture|unit|chapter)\s*(?:number\s*)?({english_ordinal}|[0-9]+)\b\s*[:：\-–—.]?\s*(.*)$",
        value,
        re.I,
    )
    if english:
        return (english.group(2).strip() or value)

    # Headings commonly produced by OCR: "Lesson 1" / "الدرس 1" without a title.
    if re.match(r"^\s*(?:Lesson|Lecture|Unit|Chapter)\s+[0-9]+\s*$", value, re.I):
        return value
    if re.match(r"^\s*(?:الدرس|المحاضرة|الوحدة|الفصل)\s+[0-9٠-٩]+\s*$", normalized):
        return value
    return None


def split_lessons(text: str) -> list[tuple[str, str]]:
    """Split one uploaded document into independent lessons.

    Priority: explicit [[LESSON:...]] markers/bookmarks, then Arabic/English
    lesson headings. If no boundary exists, keep the document as one lesson.
    """
    text = clean_text(text)
    if not text:
        return []

    lines = text.splitlines()
    starts: list[tuple[int, str]] = []
    for i, line in enumerate(lines):
        marker = LESSON_MARKER_RE.fullmatch(line.strip())
        if marker:
            title = marker.group(1).strip()
            starts.append((i, title or f"الدرس {len(starts) + 1}"))
            continue
        title = _lesson_heading(line)
        if title:
            starts.append((i, title))

    # Remove duplicate marker/heading pairs at the same location.
    deduped: list[tuple[int, str]] = []
    for start in starts:
        if deduped and start[0] <= deduped[-1][0] + 1:
            if len(start[1]) > len(deduped[-1][1]):
                deduped[-1] = start
        else:
            deduped.append(start)
    starts = deduped

    if not starts:
        return [("الدرس الكامل", text)]

    lessons = []
    for n, (start, title) in enumerate(starts):
        end = starts[n + 1][0] if n + 1 < len(starts) else len(lines)
        body_lines = [line for line in lines[start:end] if not LESSON_MARKER_RE.fullmatch(line.strip())]
        body = "\n".join(body_lines).strip()
        if body and len(body) >= 20:
            lessons.append((title or f"الدرس {n + 1}", body))

    if starts[0][0] > 0 and lessons:
        intro = "\n".join(lines[:starts[0][0]]).strip()
        if intro and len(intro) < 3000:
            lessons[0] = (lessons[0][0], intro + "\n" + lessons[0][1])

    return lessons or [("الدرس الكامل", text)]


def chunk_text(text: str, max_chars: int = 10000) -> list[str]:
    text = clean_text(text)
    return [text[i:i + max_chars] for i in range(0, len(text), max_chars)] if text else []
